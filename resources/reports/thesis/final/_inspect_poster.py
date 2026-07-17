"""Inspect the poster pptx and the reference image."""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

PPTX = "BM.KHCN.04.26 POSTER-PUS2026-A32.pptx"
IMG = "ChatGPT Image Jun 5, 2026, 03_48_20 PM.png"

def emu_to_cm(v):
    try:
        return round(Emu(v).cm, 2)
    except Exception:
        return v

# --- image info ---
try:
    im = Image.open(IMG)
    print(f"IMAGE: size={im.size} mode={im.mode}")
except Exception as e:
    print("IMAGE error:", e)

# --- pptx info ---
prs = Presentation(PPTX)
print(f"\nPPTX slide size: {emu_to_cm(prs.slide_width)} x {emu_to_cm(prs.slide_height)} cm")
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n===== SLIDE {idx} (1-based {idx+1}) layout='{slide.slide_layout.name}' =====")
    for sh in slide.shapes:
        txt = ""
        if sh.has_text_frame:
            txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text)
        print(f"  [{sh.shape_id}] {sh.shape_type} name='{sh.name}' "
              f"pos=({emu_to_cm(sh.left)},{emu_to_cm(sh.top)}) "
              f"size=({emu_to_cm(sh.width)}x{emu_to_cm(sh.height)})"
              + (f" TEXT='{txt[:90]}'" if txt else ""))
